"""Tests for the B19 per-document IDP processing pipeline.

The digital path, routing, persistence, idempotency, config-reader and endpoint tests
run anywhere. The scanned/OCR path needs PaddleOCR and is covered by a separate, skipped
test (runs on machines/CI where paddle is installed).
"""

import io
from uuid import uuid4

import pytest

from agentcore.services.idp import pipeline
from agentcore.services.idp.agent_config import resolve_pipeline_config


@pytest.fixture
def anyio_backend():
    return "asyncio"


class _MockStorage:
    """In-memory storage (the deployed default is Azure Blob, unreachable in tests)."""

    def __init__(self):
        self.store: dict[tuple[str, str], bytes] = {}

    async def save_file(self, agent_id: str, file_name: str, data: bytes) -> None:
        self.store[(agent_id, file_name)] = data

    async def get_file(self, agent_id: str, file_name: str) -> bytes:
        key = (agent_id, file_name)
        if key not in self.store:
            raise FileNotFoundError(file_name)
        return self.store[key]


@pytest.fixture(autouse=True)
def mock_storage(monkeypatch):
    s = _MockStorage()
    monkeypatch.setattr(pipeline, "get_storage_service", lambda: s)
    return s


# ─────────────────────────────── helpers ───────────────────────────────
def _node(display_name: str, fields: dict) -> dict:
    return {
        "id": display_name.replace(" ", ""),
        "type": "genericNode",
        "data": {
            "node": {
                "display_name": display_name,
                "template": {k: {"value": v} for k, v in fields.items()},
            }
        },
    }


def _graph(model_uuid: str | None, mode: str = "dynamic_prompt", prompt: str = "Extract fields", extra_nodes=None):
    nodes = [_node("AI Field Extractor", {"extraction_mode": mode, "prompt": prompt, "config_name": ""})]
    if model_uuid is not None:
        nodes.append(_node("Large Language Model", {"registry_model": f"GPT | gpt-4o | {model_uuid}"}))
    if extra_nodes:
        nodes.extend(extra_nodes)
    return {"nodes": nodes, "edges": []}


def _digital_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(50, 770, "Invoice Number INV-001 Date 2026 Total Amount 100")
    y = 748
    for r in range(6):  # plenty of on-page words so it classifies as digital
        c.drawString(50, y, " ".join(f"alpha{r}{i}" for i in range(8)))
        y -= 18
    c.showPage()
    c.save()
    return buf.getvalue()


async def _setup_document(session, *, graph, file_bytes=None, file_type="pdf", default_rule_action="pending_review", extra=None):
    """Create base Agent (with graph), IdpAgent, and a stored IdpDocument. Returns (agent_id, doc_id)."""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.database.models.idp.documents import IdpDocument

    agent_id = uuid4()
    doc_id = uuid4()

    base_agent = Agent(id=agent_id, name="IDP Pipeline Test Agent", data=graph)
    session.add(base_agent)
    await session.flush()

    idp_agent = IdpAgent(
        id=agent_id, agent_id=agent_id, extraction_mode="dynamic_prompting",
        default_rule_action=default_rule_action, extra=extra,
    )
    session.add(idp_agent)
    await session.flush()

    storage_filename = f"idp_{doc_id}.{file_type}"
    if file_bytes is not None:
        # pipeline.get_storage_service is monkeypatched to the in-memory mock
        await pipeline.get_storage_service().save_file(agent_id=str(agent_id), file_name=storage_filename, data=file_bytes)

    doc = IdpDocument(
        id=doc_id, agent_id=agent_id, original_filename=f"test.{file_type}",
        file_path=f"{agent_id}/{storage_filename}", file_type=file_type,
        file_size_bytes=len(file_bytes or b""), source="upload", status="queued",
    )
    session.add(doc)
    await session.commit()
    return agent_id, doc_id


async def _cleanup(agent_id, doc_id):
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    async with session_scope() as session:
        for job in (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).all():
            await session.delete(job)
        doc = await session.get(IdpDocument, doc_id)
        if doc:
            await session.delete(doc)
        idp = await session.get(IdpAgent, agent_id)
        if idp:
            await session.delete(idp)
        ag = await session.get(Agent, agent_id)
        if ag:
            await session.delete(ag)
        await session.commit()


def _patch_extraction(monkeypatch, *, result=None, raises=False):
    async def fake_build_llm(session, model_id):
        return object()

    async def fake_extract_dynamic(ocr_text, prompt, llm_model):
        if raises:
            raise RuntimeError("boom from LLM")
        return result if result is not None else {
            "headers": {"invoice_number": {"value": "INV-001", "confidence": 0.95, "reasoning": "top"}},
            "line_items": [],
        }

    monkeypatch.setattr(pipeline, "_build_llm", fake_build_llm)
    monkeypatch.setattr(pipeline, "extract_dynamic", fake_extract_dynamic)


# ─────────────────────────────── tests ───────────────────────────────
@pytest.mark.anyio
async def test_pipeline_digital_happy_path(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob, IdpExtractedHeader
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())

    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review")
            assert doc.processing_started_at is not None and doc.processing_completed_at is not None
            assert doc.overall_confidence is not None

            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert job.status == "completed"
            assert job.extraction_mode_used == "dynamic_prompt"
            assert isinstance(job.log, list) and job.log
            steps = [e["step"] for e in job.log]
            for s in ("load", "detect", "native", "pages", "extract", "route", "finalize"):
                assert s in steps  # digital doc -> native text (no OCR)
            assert "ocr" not in steps  # a digital PDF must NOT be re-OCR'd

            headers = (await session.exec(select(IdpExtractedHeader).where(IdpExtractedHeader.job_id == job.id))).all()
            assert any(h.field_name == "invoice_number" and h.extracted_value == "INV-001" for h in headers)
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_routing_auto_approve(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    # high-confidence result + auto_approve policy => auto_approved
    _patch_extraction(monkeypatch, result={
        "headers": {"f": {"value": "INV-001", "confidence": 0.99}}, "line_items": [],
    })
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), default_rule_action="auto_approve"
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "auto_approved"
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_routing_empty_goes_to_review(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    # zero fields => confidence 0 => pending_review even under auto_approve
    _patch_extraction(monkeypatch, result={"headers": {}, "line_items": []})
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), default_rule_action="auto_approve"
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "pending_review"
    finally:
        await _cleanup(agent_id, doc_id)


async def _add_rule(session, idp_agent_id, **kw):
    from agentcore.services.database.models.idp.config import IdpAgentRule

    defaults = dict(
        idp_agent_id=idp_agent_id, rule_group=1, condition_type="confidence_overall",
        field_name=None, operator=">=", value="0.8", action="auto_approve", display_order=0,
    )
    defaults.update(kw)
    rule = IdpAgentRule(**defaults)
    session.add(rule)
    await session.commit()
    return rule.id


async def _delete_rules(idp_agent_id):
    from agentcore.services.database.models.idp.config import IdpAgentRule
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    async with session_scope() as session:
        for r in (await session.exec(select(IdpAgentRule).where(IdpAgentRule.idp_agent_id == idp_agent_id))).all():
            await session.delete(r)
        await session.commit()


@pytest.mark.anyio
async def test_pipeline_rules_passing_rule_auto_approves(monkeypatch):
    """A passing rule (action=auto_approve) overrides a pending_review default."""
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    _patch_extraction(monkeypatch, result={"headers": {"total": {"value": "100", "confidence": 0.99}}, "line_items": []})
    async with session_scope() as session:
        # default is pending_review -> only the RULE can flip it to auto_approved
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), default_rule_action="pending_review"
        )
        await _add_rule(session, agent_id, condition_type="confidence_overall", operator=">=", value="0.8", action="auto_approve")
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "auto_approved"  # rule passed -> auto-approved despite default
    finally:
        await _delete_rules(agent_id)
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_rules_failing_rule_goes_to_review(monkeypatch):
    """A failing auto-approve rule leaves the doc in pending_review (no group matched -> default)."""
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    _patch_extraction(monkeypatch, result={"headers": {"total": {"value": "100", "confidence": 0.99}}, "line_items": []})
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), default_rule_action="pending_review"
        )
        # total=100 is NOT > 1000 -> rule fails -> falls back to default (pending_review)
        await _add_rule(
            session, agent_id, condition_type="field_value_numeric",
            field_name="total", operator=">", value="1000", action="auto_approve",
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "pending_review"
    finally:
        await _delete_rules(agent_id)
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_failure_path(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, raises=True)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "failed" and doc.failed_at is not None
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert job.status == "failed" and job.error_message
            assert job.log[-1]["status"] == "error"
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_idempotency_refuses_running(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
        # pre-existing running job (different id)
        running = IdpProcessingJob(document_id=doc_id, agent_id=agent_id, status="running")
        session.add(running)
        # our queued job
        ours = IdpProcessingJob(document_id=doc_id, agent_id=agent_id, status="queued")
        session.add(ours)
        await session.commit()
        our_job_id = ours.id
    try:
        await pipeline.process_document(doc_id, our_job_id)
        async with session_scope() as session:
            our = await session.get(IdpProcessingJob, our_job_id)
            assert our.status == "failed" and "already running" in (our.error_message or "")
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "queued"  # untouched
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_missing_model_in_graph(monkeypatch):
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(None), file_bytes=_digital_pdf()  # no LLM node
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "failed"
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert "model" in (job.error_message or "").lower()
    finally:
        await _cleanup(agent_id, doc_id)


def test_selected_pages_never_empty():
    """Page selection must never drop the whole document; out-of-range falls back to all."""
    from types import SimpleNamespace as NS

    def cfg(**k):
        base = dict(page_selection_mode="all", first_n_pages=3, page_range_start=1, page_range_end=5)
        base.update(k)
        return NS(**base)

    sel = pipeline._selected_pages
    assert sel(5, cfg(page_selection_mode="all")) == {1, 2, 3, 4, 5}
    assert sel(12, cfg(page_selection_mode="first_n", first_n_pages=3)) == {1, 2, 3}
    assert sel(8, cfg(page_selection_mode="last_n", first_n_pages=2)) == {7, 8}
    # range entirely past the doc -> all pages (was: empty set, dropping everything)
    assert sel(5, cfg(page_selection_mode="range", page_range_start=8, page_range_end=10)) == {1, 2, 3, 4, 5}
    # normal in-range
    assert sel(10, cfg(page_selection_mode="range", page_range_start=2, page_range_end=4)) == {2, 3, 4}


@pytest.mark.anyio
async def test_pipeline_extraction_error_is_fatal(monkeypatch):
    """An extractor error with zero fields is fatal (consistent across modes), not a clean review."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    # extract_dynamic swallows LLM errors into {"error":...}; with zero fields -> must fail
    _patch_extraction(monkeypatch, result={"headers": {}, "line_items": [], "error": "LLM timeout"})
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "failed"
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert "LLM timeout" in (job.error_message or "")
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_multisheet_xlsx_keeps_all_sheets(monkeypatch, mock_storage):
    """A multi-sheet workbook must process ALL sheets (regression: total_pages truncated to sheet 1)."""
    from openpyxl import Workbook
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    wb = Workbook()
    wb.active["A1"] = "alpha sheet one"
    s2 = wb.create_sheet("Second")
    s2["A1"] = "beta sheet two"
    buf = io.BytesIO()
    wb.save(buf)
    xlsx = buf.getvalue()

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=xlsx, file_type="xlsx"
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.page_count == 2  # both sheets counted (was 1 -> sheet 2 dropped)
        saved = [v for (a, f), v in mock_storage.store.items() if f.endswith("ocr_text.txt")]
        assert saved and b"beta sheet two" in saved[0]  # sheet 2 content survived
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_enqueue_reaps_queue(monkeypatch):
    """enqueue_document must clean up its JobQueueService entry once the job finishes (no leak)."""
    import asyncio as _asyncio
    from agentcore.services.deps import session_scope, get_queue_service

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        async with session_scope() as session:
            job_id = await pipeline.enqueue_document(session, doc_id)
        qs = get_queue_service()
        assert str(job_id) in qs._queues  # registered while running
        for _ in range(80):  # wait for job task + monitor to reap
            if str(job_id) not in qs._queues:
                break
            await _asyncio.sleep(0.1)
        assert str(job_id) not in qs._queues, "queue entry leaked after job completion"
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_agent_config_reader():
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.deps import session_scope

    model_uuid = str(uuid4())
    extra_nodes = [
        _node("Scan Corrector", {"fix_skew": False, "fix_rotation": True, "allowed_angles": "90,270"}),
        _node("Page Selector", {"selection_mode": "first_n", "first_n_pages": 2}),
        _node("Document Type Detector", {"allowed_extensions": "pdf,png", "min_text_length": 40}),
    ]
    graph = _graph(model_uuid, mode="dynamic_prompt", prompt="Read it", extra_nodes=extra_nodes)
    idp_agent = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="dynamic_prompting", default_rule_action="auto_approve")
    base_agent = Agent(id=uuid4(), name="x", data=graph)

    async with session_scope() as session:
        cfg = await resolve_pipeline_config(session, idp_agent, base_agent)

    assert cfg.model_id == model_uuid
    assert cfg.extraction_mode == "dynamic_prompt"
    assert cfg.prompt == "Read it"
    assert cfg.fix_skew is False and cfg.fix_rotation is True
    assert cfg.allowed_angles == [90, 270]
    assert cfg.page_selection_mode == "first_n" and cfg.first_n_pages == 2
    assert cfg.allowed_extensions == {"pdf", "png"} and cfg.min_text_length == 40
    assert cfg.default_rule_action == "auto_approve"


@pytest.mark.anyio
async def test_process_endpoint(monkeypatch):
    from fastapi.testclient import TestClient
    from agentcore.main import create_app
    from agentcore.services.auth.utils import get_current_active_user
    from agentcore.api.idp import idp_rbac, _idp_submit_rbac
    from agentcore.api.idp import documents as documents_api
    from agentcore.services.database.models.idp.documents import IdpProcessingJob
    from agentcore.services.deps import session_scope

    def mock_user():
        from types import SimpleNamespace
        return SimpleNamespace(id=uuid4(), role="super_admin")

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=b"%PDF-1.4 test")

    fixed_job = uuid4()

    async def fake_enqueue(session, document_id):
        return fixed_job

    monkeypatch.setattr(documents_api, "enqueue_document", fake_enqueue)

    app = create_app()
    app.dependency_overrides[get_current_active_user] = mock_user
    app.dependency_overrides[idp_rbac] = mock_user
    app.dependency_overrides[_idp_submit_rbac] = mock_user
    client = TestClient(app)
    try:
        r = client.post(f"/api/v1/idp/documents/{doc_id}/process")
        assert r.status_code == 202, r.text
        assert r.json()["job_id"] == str(fixed_job)

        # a running job now exists -> 409
        async with session_scope() as session:
            session.add(IdpProcessingJob(document_id=doc_id, agent_id=agent_id, status="running"))
            await session.commit()
        r2 = client.post(f"/api/v1/idp/documents/{doc_id}/process")
        assert r2.status_code == 409, r2.text
    finally:
        app.dependency_overrides.clear()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_persists_flat_line_items(monkeypatch):
    """Flat LLM line-item shape (no `columns`) must still persist via the normalizer."""
    from agentcore.services.database.models.idp.documents import IdpProcessingJob, IdpExtractedLineItem
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, result={
        "headers": {"invoice_number": {"value": "INV-1", "confidence": 0.9}},
        "line_items": [
            {"Item": "Widget A", "Qty": "10", "Amount": "50.00"},
            {"Item": "Gadget B", "Qty": "3", "Amount": "60.00"},
        ],
    })
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            lines = (await session.exec(select(IdpExtractedLineItem).where(IdpExtractedLineItem.job_id == job.id))).all()
            assert len(lines) == 6  # 2 rows x 3 columns
            assert {l.column_name for l in lines} == {"Item", "Qty", "Amount"}
            assert any(l.column_name == "Item" and l.extracted_value == "Widget A" for l in lines)
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_document_file_endpoint(monkeypatch):
    """GET /documents/{id}/file streams the stored bytes for the viewer."""
    from fastapi.testclient import TestClient
    from agentcore.main import create_app
    from agentcore.services.auth.utils import get_current_active_user
    from agentcore.api.idp import idp_rbac
    from agentcore.services.deps import session_scope, get_storage_service

    def mock_user():
        from types import SimpleNamespace
        return SimpleNamespace(id=uuid4(), role="root")  # root bypasses org scoping

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=b"%PDF-1.4 hello")

    class FileStore:
        async def get_file(self, agent_id, file_name):
            return b"%PDF-1.4 hello"

    app = create_app()
    app.dependency_overrides[get_current_active_user] = mock_user
    app.dependency_overrides[idp_rbac] = mock_user
    app.dependency_overrides[get_storage_service] = lambda: FileStore()
    client = TestClient(app)
    try:
        r = client.get(f"/api/v1/idp/documents/{doc_id}/file")
        assert r.status_code == 200, r.text
        assert r.content == b"%PDF-1.4 hello"
        assert "application/pdf" in r.headers.get("content-type", "")
    finally:
        app.dependency_overrides.clear()
        await _cleanup(agent_id, doc_id)


def _digital_pdf_pages(n: int) -> bytes:
    """An n-page born-digital PDF (each page word-dense so it classifies digital)."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for p in range(n):
        c.drawString(50, 770, f"Invoice Number INV-{p:03d} Total Amount {100 + p}")
        y = 748
        for r in range(6):
            c.drawString(50, y, " ".join(f"word{p}{r}{i}" for i in range(8)))
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def test_long_doc_build_sections_splits_on_headings():
    from agentcore.services.idp import long_doc

    tokens = [
        {"text": "introduction paragraph one", "page_number": 1},
        {"text": "SUMMARY", "page_number": 2},
        {"text": "more body text on page three", "page_number": 3},
    ]
    sections = long_doc.build_sections(tokens)
    # page 1 -> "Document"; page 2 opens with heading SUMMARY -> new section (extends to p3)
    assert len(sections) == 2
    assert sections[0]["page_start"] == 1 and sections[0]["page_end"] == 1
    assert sections[1]["title"] == "SUMMARY" and sections[1]["page_start"] == 2 and sections[1]["page_end"] == 3


def test_long_doc_merge_chunk_extractions():
    from agentcore.services.idp import long_doc

    results = [
        {"headers": {"a": {"value": "1", "confidence": 0.5}},
         "line_items": [{"row_index": 0, "columns": [{"column_name": "c", "value": "x"}]}]},
        {"headers": {"a": {"value": "2", "confidence": 0.9}, "b": {"value": "3", "confidence": 0.7}},
         "line_items": [{"row_index": 0, "columns": [{"column_name": "c", "value": "y"}]}]},
    ]
    merged = long_doc.merge_chunk_extractions(results)
    assert merged["headers"]["a"]["value"] == "2"  # higher-confidence value wins
    assert merged["headers"]["b"]["value"] == "3"
    assert [r["row_index"] for r in merged["line_items"]] == [0, 1]  # continuous re-index


@pytest.mark.anyio
async def test_pipeline_long_doc_chunks_and_persists_sections(monkeypatch):
    """A many-page digital doc triggers chunking + writes idp_document_sections."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob, IdpDocumentSection
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)  # mock LLM returns the same header for every chunk
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf_pages(10)  # >8 pages -> long
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review")
            assert doc.page_count == 10
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert "long_doc" in [e["step"] for e in job.log]
            sections = (await session.exec(select(IdpDocumentSection).where(IdpDocumentSection.document_id == doc_id))).all()
            assert len(sections) >= 1  # at least one section persisted
    finally:
        async with session_scope() as session:
            for s in (await session.exec(select(IdpDocumentSection).where(IdpDocumentSection.document_id == doc_id))).all():
                await session.delete(s)
            await session.commit()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_entity_linking_links_cross_page_value():
    """A header value present on >= 2 pages becomes one idp_entity_link; single-page values don't."""
    from agentcore.services.idp.entity_linking import link_entities

    class FakeSession:
        def __init__(self):
            self.added = []

        def add(self, obj):
            self.added.append(obj)

        async def commit(self):
            pass

    tokens = [
        {"text": "Vendor", "page_number": 1}, {"text": "ACME", "page_number": 1}, {"text": "CORP", "page_number": 1},
        {"text": "continued", "page_number": 2}, {"text": "ACME", "page_number": 2}, {"text": "CORP", "page_number": 2},
        {"text": "unique", "page_number": 2},
    ]
    extracted = {
        "headers": {
            "vendor": {"value": "ACME CORP", "confidence": 0.9},
            "page2only": {"value": "unique", "confidence": 0.8},
        },
        "line_items": [],
    }
    sess = FakeSession()
    n = await link_entities(sess, uuid4(), extracted, tokens)
    assert n == 1  # only the cross-page vendor; "unique" is page-2-only
    link = sess.added[0]
    assert link.entity_type == "vendor" and link.occurrences["pages"] == [1, 2]


def _two_page_pdf_with_repeat() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for p in range(2):
        c.drawString(50, 770, "Vendor ACME CORP Invoice")
        y = 748
        for r in range(6):
            c.drawString(50, y, " ".join(f"w{p}{r}{i}" for i in range(8)))
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


@pytest.mark.anyio
async def test_pipeline_writes_entity_links(monkeypatch):
    """End-to-end: a repeated header value across 2 pages writes idp_entity_links."""
    from agentcore.services.database.models.idp.documents import IdpEntityLink
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, result={"headers": {"vendor": {"value": "ACME CORP", "confidence": 0.9}}, "line_items": []})
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_two_page_pdf_with_repeat())
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            links = (await session.exec(select(IdpEntityLink).where(IdpEntityLink.document_id == doc_id))).all()
            assert any(l.entity_type == "vendor" and l.occurrences.get("pages") == [1, 2] for l in links)
    finally:
        async with session_scope() as session:
            for l in (await session.exec(select(IdpEntityLink).where(IdpEntityLink.document_id == doc_id))).all():
                await session.delete(l)
            await session.commit()
        await _cleanup(agent_id, doc_id)


# ───────────────────── Codex-review fixes (2026-06-11) ─────────────────────
def test_as_bool_coercion():
    """agent_config._as_bool must interpret string toggles, not use truthiness."""
    from agentcore.services.idp.agent_config import _as_bool

    assert _as_bool("false", True) is False  # the bug: bool("false") was True
    assert _as_bool("true", False) is True
    assert _as_bool("0", True) is False and _as_bool("1", False) is True
    assert _as_bool(True, False) is True and _as_bool(False, True) is False
    assert _as_bool(None, True) is True  # missing -> default
    assert _as_bool(0, True) is False and _as_bool(2, False) is True


def test_chunks_for_extraction_subsplits_huge_page():
    """A single page larger than max_tokens must be fixed-token-split, not sent whole."""
    from agentcore.services.idp import long_doc

    tokens = [{"text": "word " * 2000, "page_number": 1}]  # ~10k chars on ONE page
    sections = [{"title": "Document", "page_start": 1, "page_end": 1, "order_index": 0, "parent_idx": None}]
    chunks = long_doc.chunks_for_extraction(tokens, sections, max_tokens=1000)  # limit ~4000 chars
    assert len(chunks) > 1  # the oversized page was split into multiple chunks
    assert all(c.strip() for c in chunks)


@pytest.mark.anyio
async def test_hook_split_skips_child_document():
    """A split-child (parent_document_id set) must never be split again."""
    from types import SimpleNamespace
    from agentcore.services.idp.pipeline import _hook_split, FlowLog

    cfg = SimpleNamespace(multi_doc_split=True)
    child = SimpleNamespace(parent_document_id=uuid4())
    flow = FlowLog(uuid4(), "child.pdf")
    result = await _hook_split(None, None, child, [], {}, cfg, flow)
    assert result is None
    assert any("child document" in e["detail"] for e in flow.events)


@pytest.mark.anyio
async def test_pipeline_long_doc_total_failure_is_fatal(monkeypatch):
    """If every long-doc chunk extraction fails, the document must end 'failed', not review."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, raises=True)  # LLM raises for every chunk
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf_pages(10))
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "failed"
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert job.status == "failed" and "chunk" in (job.error_message or "").lower()
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_idp_documents_allows_split_status():
    """The idp_b22 migration must allow status='split' (split-parent container)."""
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=b"%PDF-1.4 x")
    try:
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            doc.status = "split"
            session.add(doc)
            await session.commit()  # must NOT violate ck_idp_documents_status
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status == "split"
    finally:
        await _cleanup(agent_id, doc_id)


# ───────── differentiator HOOK WIRING (pipeline fires them + persists) ─────────
@pytest.mark.anyio
async def test_pipeline_fires_classification_hook(monkeypatch):
    """With classify_enabled, the pipeline calls classification and an idp_document_classifications row lands."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpDocumentClassification
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)

    async def fake_classify(session, document_id, merged_text, org_id, *, auto_select, threshold, doc_types=None):
        session.add(IdpDocumentClassification(
            document_id=document_id, predicted_type="Invoice", confidence=0.95, is_selected=False,
        ))
        await session.commit()
        return {"predicted_type": "Invoice", "confidence": 0.95, "candidates": None, "selected_config_id": None}

    monkeypatch.setattr("agentcore.services.idp.classification.classify_and_persist", fake_classify)

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(),
            extra={"classify_enabled": True, "classify_auto_select": False},
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            rows = (await session.exec(select(IdpDocumentClassification).where(IdpDocumentClassification.document_id == doc_id))).all()
            assert len(rows) == 1 and rows[0].predicted_type == "Invoice"
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review")  # pipeline still completes
    finally:
        async with session_scope() as session:
            for r in (await session.exec(select(IdpDocumentClassification).where(IdpDocumentClassification.document_id == doc_id))).all():
                await session.delete(r)
            await session.commit()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_fires_detection_hook(monkeypatch):
    """With detect_enabled, the pipeline calls detection and an idp_detected_elements row lands."""
    from agentcore.services.database.models.idp.documents import IdpDetectedElement
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)

    async def fake_detect(session, document_id, file_bytes, file_type, enabled_types):
        session.add(IdpDetectedElement(
            document_id=document_id, element_type="signature", page_number=1,
            bounding_box={"x_min": 0.0, "y_min": 0.0, "x_max": 1.0, "y_max": 1.0}, confidence=0.9,
        ))
        await session.commit()
        return ["ok"]

    monkeypatch.setattr("agentcore.services.idp.visual_detection.detect_and_persist", fake_detect)

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), extra={"detect_enabled": True},
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            rows = (await session.exec(select(IdpDetectedElement).where(IdpDetectedElement.document_id == doc_id))).all()
            assert len(rows) == 1 and rows[0].element_type == "signature"
    finally:
        async with session_scope() as session:
            for r in (await session.exec(select(IdpDetectedElement).where(IdpDetectedElement.document_id == doc_id))).all():
                await session.delete(r)
            await session.commit()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_split_early_returns_and_enqueues_children(monkeypatch):
    """With multi_doc_split + >=2 boundaries, parent becomes 'split', children are enqueued, no extraction."""
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpExtractedHeader, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)
    enqueued: list = []

    def fake_boundaries(tokens, page_status):
        return [(0, 0), (1, 1)]

    async def fake_materialize(session, storage, parent_doc, boundaries):
        ids = []
        for i, _ in enumerate(boundaries):
            child = IdpDocument(
                agent_id=parent_doc.agent_id, parent_document_id=parent_doc.id,
                original_filename=f"child{i}.pdf", file_path=f"{parent_doc.agent_id}/child{i}.pdf",
                file_type="pdf", file_size_bytes=10, source="upload", status="queued",
                extra={"multi_doc_split": False},
            )
            session.add(child)
            await session.flush()
            ids.append(child.id)
        await session.commit()
        return ids

    async def fake_enqueue(session, document_id):
        enqueued.append(document_id)
        return uuid4()

    monkeypatch.setattr("agentcore.services.idp.splitting.detect_document_boundaries", fake_boundaries)
    monkeypatch.setattr("agentcore.services.idp.splitting.materialize_children", fake_materialize)
    monkeypatch.setattr(pipeline, "enqueue_document", fake_enqueue)

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
        idp = await session.get(IdpAgent, agent_id)
        idp.multi_doc_split = True  # split is read from the idp_agent column
        session.add(idp)
        await session.commit()
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            parent = await session.get(IdpDocument, doc_id)
            assert parent.status == "split"  # fix #1: 'split' is an allowed status
            children = (await session.exec(select(IdpDocument).where(IdpDocument.parent_document_id == doc_id))).all()
            assert len(children) == 2
            assert len(enqueued) == 2  # both children enqueued
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert "split" in [e["step"] for e in job.log]
            # no extraction happened on the parent (it's a container)
            heads = (await session.exec(select(IdpExtractedHeader).where(IdpExtractedHeader.job_id == job.id))).all()
            assert heads == []
    finally:
        async with session_scope() as session:
            for c in (await session.exec(select(IdpDocument).where(IdpDocument.parent_document_id == doc_id))).all():
                await session.delete(c)
            await session.commit()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_flow_log_is_detailed(monkeypatch):
    """The per-document flow log records the full cycle: input, config, OCR/native, what the LLM gave."""
    from agentcore.services.database.models.idp.documents import IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, result={
        "headers": {"invoice_number": {"value": "INV-9", "confidence": 0.9}, "total": {"value": "100", "confidence": 0.9}},
        "line_items": [],
    })
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            by_step = {e["step"]: e["detail"] for e in job.log}
            # input received
            assert "input received" in by_step["load"]
            # the agent config (what runs) is logged
            assert "config" in by_step and "extraction=" in by_step["config"] and "features=" in by_step["config"]
            # OCR/native output size is logged
            assert "native" in by_step and "chars" in by_step["native"]
            # merged text size + location
            assert "text" in by_step and "chars" in by_step["text"] and "ocr_output" in by_step["text"]
            # what the LLM gave (field names)
            assert "invoice_number" in by_step["extract"] and "header field" in by_step["extract"]
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_classify_failure_does_not_poison_session(monkeypatch):
    """A classification hook failure must roll back and NOT corrupt extraction/finalize."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpExtractedHeader, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)

    async def boom_classify(session, document_id, merged_text, org_id, *, auto_select, threshold):
        # The hook now runs this on an ISOLATED session, so even a DB-poisoning failure
        # (out-of-range confidence violating the CHECK on commit) must not touch the main
        # pipeline session. This is the real scenario the isolation fix protects against.
        from agentcore.services.database.models.idp.documents import IdpDocumentClassification
        session.add(IdpDocumentClassification(document_id=document_id, predicted_type="X", confidence=1.5))
        await session.commit()  # CheckViolation -> rolls back the ISOLATED session, re-raises

    monkeypatch.setattr("agentcore.services.idp.classification.classify_and_persist", boom_classify)

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(
            session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf(), extra={"classify_enabled": True},
        )
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review")  # pipeline survived the bad hook
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert job.status == "completed"
            heads = (await session.exec(select(IdpExtractedHeader).where(IdpExtractedHeader.job_id == job.id))).all()
            assert len(heads) >= 1  # extraction still persisted despite the classify failure
            assert "classify" in [e["step"] for e in job.log]  # the warn was logged
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_pipeline_split_all_enqueue_fail_fails_parent(monkeypatch):
    """If split produces children but NONE can be enqueued, the parent fails (not silently 'split')."""
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)

    def fake_boundaries(tokens, page_status):
        return [(0, 0), (1, 1)]

    async def fake_materialize(session, storage, parent_doc, boundaries):
        ids = []
        for i, _ in enumerate(boundaries):
            child = IdpDocument(
                agent_id=parent_doc.agent_id, parent_document_id=parent_doc.id,
                original_filename=f"c{i}.pdf", file_path=f"{parent_doc.agent_id}/c{i}.pdf",
                file_type="pdf", file_size_bytes=10, source="upload", status="queued",
            )
            session.add(child)
            await session.flush()
            ids.append(child.id)
        await session.commit()
        return ids

    async def always_fail_enqueue(session, document_id):
        raise RuntimeError("queue down")

    monkeypatch.setattr("agentcore.services.idp.splitting.detect_document_boundaries", fake_boundaries)
    monkeypatch.setattr("agentcore.services.idp.splitting.materialize_children", fake_materialize)
    monkeypatch.setattr(pipeline, "enqueue_document", always_fail_enqueue)

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
        idp = await session.get(IdpAgent, agent_id)
        idp.multi_doc_split = True
        session.add(idp)
        await session.commit()
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            parent = await session.get(IdpDocument, doc_id)
            assert parent.status == "failed"  # not 'split' — nothing got enqueued
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            assert job.status == "failed" and "enqueue" in (job.error_message or "").lower()
    finally:
        async with session_scope() as session:
            for c in (await session.exec(select(IdpDocument).where(IdpDocument.parent_document_id == doc_id))).all():
                await session.delete(c)
            await session.commit()
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_agent_config_reads_differentiator_nodes():
    """Classifier / Visual Detection / Math Reconcile canvas nodes drive the pipeline toggles."""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.deps import session_scope

    extra_nodes = [
        _node("Document Classifier", {"confidence_threshold": 0.6}),
        _node("Visual Element Detection", {"detect_signatures": True, "detect_qr": True, "detect_stamps": False}),
        _node("Math Reconcile", {"max_retries": 3, "tolerance": 0.05}),
    ]
    graph = _graph(str(uuid4()), extra_nodes=extra_nodes)
    idp_agent = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="dynamic_prompting", default_rule_action="pending_review")
    base_agent = Agent(id=uuid4(), name="x", data=graph)
    async with session_scope() as session:
        cfg = await resolve_pipeline_config(session, idp_agent, base_agent)

    assert cfg.classify_enabled is True and cfg.classify_threshold == 0.6
    # detect_qr maps to both qr + barcode (pyzbar emits both); stamps not supported -> ignored
    assert cfg.detect_enabled is True and cfg.detect_enabled_types == {"signature", "qr", "barcode"}
    assert cfg.math_reconcile_enabled is True
    assert cfg.math_reconcile_max_attempts == 3 and cfg.math_reconcile_tolerance == 0.05


@pytest.mark.anyio
async def test_agent_config_no_differentiator_nodes_defaults_off():
    """Without the nodes, classify/detect/math default OFF; idp_agent.extra still overrides."""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent
    from agentcore.services.deps import session_scope

    graph = _graph(str(uuid4()))  # no differentiator nodes
    idp_off = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="dynamic_prompting", default_rule_action="pending_review")
    idp_extra = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="dynamic_prompting",
                         default_rule_action="pending_review", extra={"classify_enabled": True})
    base_agent = Agent(id=uuid4(), name="x", data=graph)
    async with session_scope() as session:
        cfg_off = await resolve_pipeline_config(session, idp_off, base_agent)
        cfg_extra = await resolve_pipeline_config(session, idp_extra, base_agent)

    assert cfg_off.classify_enabled is False and cfg_off.detect_enabled is False and cfg_off.math_reconcile_enabled is False
    assert cfg_extra.classify_enabled is True  # extra fallback still works


@pytest.mark.anyio
async def test_pipeline_txt_file(monkeypatch):
    """A .txt file is accepted, read as digital native text (no OCR), and processed."""
    from agentcore.services.database.models.idp.documents import IdpDocument, IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch)
    txt = b"Invoice Number INV-TXT-1\nVendor: ACME Supplies Ltd\nSubtotal 100.00\nTax 18.00\nTotal 118.00\n"
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=txt, file_type="txt")
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review")
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            steps = [e["step"] for e in job.log]
            assert "native" in steps and "ocr" not in steps  # txt = digital native text, no OCR
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_document_log_endpoint(monkeypatch):
    """GET /documents/{id}/log returns the per-document flow log after processing."""
    from fastapi.testclient import TestClient
    from agentcore.main import create_app
    from agentcore.services.auth.utils import get_current_active_user
    from agentcore.api.idp import idp_rbac
    from agentcore.services.deps import session_scope

    def mock_user():
        from types import SimpleNamespace
        return SimpleNamespace(id=uuid4(), role="root")  # root bypasses org scoping

    _patch_extraction(monkeypatch)
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())
    try:
        await pipeline.process_document(doc_id)
        app = create_app()
        app.dependency_overrides[get_current_active_user] = mock_user
        app.dependency_overrides[idp_rbac] = mock_user
        client = TestClient(app)
        try:
            r = client.get(f"/api/v1/idp/documents/{doc_id}/log")
            assert r.status_code == 200, r.text
            body = r.json()
            steps = [e["step"] for e in body["log"]]
            assert "load" in steps and "extract" in steps and "finalize" in steps
            assert body["status"] == "completed"
            assert "input received" in body["log"][0]["detail"]
        finally:
            app.dependency_overrides.clear()
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_document_log_download_endpoint(monkeypatch):
    """GET /documents/{id}/log/download streams the flow.log artifact; 404 when absent."""
    from fastapi.testclient import TestClient
    from agentcore.main import create_app
    from agentcore.services.auth.utils import get_current_active_user
    from agentcore.api.idp import idp_rbac
    from agentcore.services.deps import session_scope, get_storage_service

    def mock_user():
        from types import SimpleNamespace
        return SimpleNamespace(id=uuid4(), role="root")  # root bypasses org scoping

    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=b"%PDF-1.4 x")

    expected_path = f"flow_logs/{doc_id}/flow.log"

    class FileStore:
        async def get_file(self, agent_id, file_name):
            if file_name == expected_path:
                return b"load | ok | input received\nextract | ok | 3 headers\nfinalize | ok"
            raise FileNotFoundError(file_name)

    class EmptyStore:
        async def get_file(self, agent_id, file_name):
            raise FileNotFoundError(file_name)

    app = create_app()
    app.dependency_overrides[get_current_active_user] = mock_user
    app.dependency_overrides[idp_rbac] = mock_user
    try:
        # artifact present → 200 + attachment + text body
        app.dependency_overrides[get_storage_service] = lambda: FileStore()
        client = TestClient(app)
        r = client.get(f"/api/v1/idp/documents/{doc_id}/log/download")
        assert r.status_code == 200, r.text
        assert b"input received" in r.content
        assert "text/plain" in r.headers.get("content-type", "")
        assert "attachment" in r.headers.get("content-disposition", "")
        assert "_processing_log.txt" in r.headers.get("content-disposition", "")

        # artifact missing → graceful 404 (never 500)
        app.dependency_overrides[get_storage_service] = lambda: EmptyStore()
        client = TestClient(app)
        r2 = client.get(f"/api/v1/idp/documents/{doc_id}/log/download")
        assert r2.status_code == 404, r2.text
    finally:
        app.dependency_overrides.clear()
        await _cleanup(agent_id, doc_id)


@pytest.mark.skipif(
    not __import__("agentcore.services.idp.ocr", fromlist=["_paddle_ocr_available"])._paddle_ocr_available,
    reason="PaddleOCR not installed (runs on M4/NVIDIA/Docker/CI)",
)
@pytest.mark.anyio
async def test_pipeline_scanned_path(monkeypatch):
    """Full OCR path with a scanned image. Skipped where paddle is absent."""
    from agentcore.services.database.models.idp.documents import IdpDocument
    from agentcore.services.deps import session_scope

    _patch_extraction(monkeypatch)
    # a tiny PNG (1x1) -> classified scanned -> real PaddleOCR
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4"
        b"\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=png)
        doc = await session.get(IdpDocument, doc_id)
        doc.file_type = "png"
        doc.file_path = doc.file_path.replace(".pdf", ".png")
        session.add(doc)
        await session.commit()
    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            doc = await session.get(IdpDocument, doc_id)
            assert doc.status in ("auto_approved", "pending_review", "failed")
    finally:
        await _cleanup(agent_id, doc_id)


def test_flow_log_io_helpers_are_bounded():
    """_clip caps text; _io_sample_rows/_io_headers cap rows/headers for the JSONB log."""
    assert pipeline._clip("x" * 5000).startswith("x" * pipeline._IO_TEXT_SAMPLE)
    assert "+" in pipeline._clip("x" * 5000)  # truncation marker
    assert pipeline._clip("short") == "short"

    # Create well above the caps so the clamp is exercised regardless of the cap values.
    extracted = {
        "headers": {f"h{i}": {"value": str(i)} for i in range(pipeline._IO_MAX_HEADERS + 20)},
        "line_items": [{"columns": [{"column_name": "a", "value": str(i)}]} for i in range(pipeline._IO_MAX_ROWS + 20)],
    }
    assert len(pipeline._io_headers(extracted)) == pipeline._IO_MAX_HEADERS
    assert len(pipeline._io_sample_rows(extracted)) == pipeline._IO_MAX_ROWS
    assert pipeline._io_sample_rows(extracted)[0] == {"a": "0"}


@pytest.mark.anyio
async def test_pipeline_flow_log_carries_io_payloads(monkeypatch):
    """The flow log must carry per-component input/output payloads for the UI to render
    the 'complete cycle' (input received, OCR/text output, what the LLM returned)."""
    from agentcore.services.database.models.idp.documents import IdpProcessingJob
    from agentcore.services.deps import session_scope
    from sqlmodel import select

    _patch_extraction(monkeypatch, result={
        "headers": {"invoice_number": {"value": "INV-IO", "confidence": 0.95}},
        "line_items": [
            {"columns": [{"column_name": "description", "value": "Row A"}, {"column_name": "amount", "value": "10"}]},
            {"columns": [{"column_name": "description", "value": "Row B"}, {"column_name": "amount", "value": "20"}]},
        ],
    })
    async with session_scope() as session:
        agent_id, doc_id = await _setup_document(session, graph=_graph(str(uuid4())), file_bytes=_digital_pdf())

    try:
        await pipeline.process_document(doc_id)
        async with session_scope() as session:
            job = (await session.exec(select(IdpProcessingJob).where(IdpProcessingJob.document_id == doc_id))).first()
            by_step = {e["step"]: e for e in job.log}

            # load -> input payload
            assert by_step["load"]["io"]["input"]["filename"]
            # native (digital) -> output text sample (string, bounded)
            nat = by_step["native"]["io"]["output"]
            assert isinstance(nat["text_sample"], str) and len(nat["text_sample"]) <= pipeline._IO_TEXT_SAMPLE + 32
            # extract -> the "what the LLM gave" payload
            ex = by_step["extract"]["io"]["output"]
            assert ex["headers"]["invoice_number"] == "INV-IO"
            assert ex["line_item_rows"] == 2
            assert ex["sample_rows"][0] == {"description": "Row A", "amount": "10"}
            # route -> decision payload
            assert by_step["route"]["io"]["output"]["decision"] in ("auto_approved", "pending_review")
    finally:
        await _cleanup(agent_id, doc_id)


@pytest.mark.anyio
async def test_named_config_resolution_prefers_real_over_template(monkeypatch):
    """When two active configs share a name (a catalogue template + a real config), resolution is
    deterministic and picks the REAL config (is_template False), never the template — and never an
    arbitrary .first() over unordered rows."""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent, IdpFieldConfiguration
    from agentcore.services.deps import session_scope

    name = f"ResolveTest-{uuid4().hex[:8]}"
    template_id, real_id = uuid4(), uuid4()
    async with session_scope() as session:
        # Insert the TEMPLATE first (older) ...
        session.add(IdpFieldConfiguration(id=template_id, name=name, is_active=True, is_template=True))
        await session.commit()
    async with session_scope() as session:
        # ... then the REAL config (newer, non-template).
        session.add(IdpFieldConfiguration(id=real_id, name=name, is_active=True, is_template=False))
        await session.commit()

    try:
        graph = {
            "nodes": [
                _node("AI Field Extractor", {"extraction_mode": "field_configuration", "config_name": name, "prompt": ""}),
                _node("Large Language Model", {"registry_model": f"GPT | gpt-4o | {uuid4()}"}),
            ],
            "edges": [],
        }
        idp_agent = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="named_config")
        base_agent = Agent(id=uuid4(), name="x", data=graph)  # org_id None -> fallback ranking path

        async with session_scope() as session:
            cfg = await resolve_pipeline_config(session, idp_agent, base_agent)
        assert cfg.field_config_id == real_id  # real config wins over the template
    finally:
        async with session_scope() as session:
            for cid in (template_id, real_id):
                row = await session.get(IdpFieldConfiguration, cid)
                if row:
                    await session.delete(row)
            await session.commit()


@pytest.mark.anyio
async def test_named_config_resolution_never_leaks_another_orgs_config():
    """Tenant safety: an agent with no org must NOT resolve to ANOTHER org's private config of the
    same name — only global (org_id NULL) configs are eligible in the fallback."""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent, IdpFieldConfiguration
    from agentcore.services.database.models.organization.model import Organization
    from agentcore.services.database.models.user.model import User
    from agentcore.services.deps import session_scope

    name = f"TenantTest-{uuid4().hex[:8]}"
    user_id, org_id = uuid4(), uuid4()
    global_id, foreign_id = uuid4(), uuid4()
    async with session_scope() as session:
        session.add(User(id=user_id, username=f"tt_{user_id.hex[:8]}@t.com", email=f"tt_{user_id.hex[:8]}@t.com",
                         password="x", is_active=True, is_superuser=False, role="developer"))
        await session.flush()
        session.add(Organization(id=org_id, name=f"TenantOrg-{org_id.hex[:8]}", owner_user_id=user_id, created_by=user_id))
        await session.flush()
        # Org B's PRIVATE config (newer) + a GLOBAL config (older), same name.
        session.add(IdpFieldConfiguration(id=global_id, name=name, is_active=True, is_template=False, org_id=None))
        await session.flush()
        session.add(IdpFieldConfiguration(id=foreign_id, name=name, is_active=True, is_template=False, org_id=org_id))
        await session.commit()

    try:
        graph = {
            "nodes": [
                _node("AI Field Extractor", {"extraction_mode": "field_configuration", "config_name": name, "prompt": ""}),
                _node("Large Language Model", {"registry_model": f"GPT | gpt-4o | {uuid4()}"}),
            ],
            "edges": [],
        }
        idp_agent = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="named_config")
        base_agent = Agent(id=uuid4(), name="x", data=graph)  # NO org -> must not see org B's config

        async with session_scope() as session:
            cfg = await resolve_pipeline_config(session, idp_agent, base_agent)
        assert cfg.field_config_id == global_id        # the global config, ...
        assert cfg.field_config_id != foreign_id       # ... never the other org's private config
    finally:
        async with session_scope() as session:
            for cid in (global_id, foreign_id):
                row = await session.get(IdpFieldConfiguration, cid)
                if row:
                    await session.delete(row)
            org = await session.get(Organization, org_id)
            if org:
                await session.delete(org)
            usr = await session.get(User, user_id)
            if usr:
                await session.delete(usr)
            await session.commit()


def test_resolve_model_id_scans_all_llm_nodes():
    """A graph with several 'Large Language Model' nodes (the universal pipeline has 2-3, one per
    consumer) must resolve the model from ANY node that carries one — not only the first node,
    which may be empty. (Universal-agent blocker fix.)"""
    from agentcore.services.idp.agent_config import _resolve_model_id_from_graph

    mid = str(uuid4())
    data = {
        "nodes": [
            _node("Large Language Model", {"registry_model": ""}),                      # first: empty
            _node("Large Language Model", {"registry_model": f"Groq | llama-3.3 | {mid}"}),  # second: filled
        ],
        "edges": [],
    }
    assert _resolve_model_id_from_graph(data) == mid
    # No LLM node carries a model -> None.
    empty = {"nodes": [_node("Large Language Model", {"registry_model": ""})], "edges": []}
    assert _resolve_model_id_from_graph(empty) is None
    # A malformed (non-UUID) third part is skipped, not returned.
    bad = {"nodes": [_node("Large Language Model", {"registry_model": "G | m | not-a-uuid"})], "edges": []}
    assert _resolve_model_id_from_graph(bad) is None
    # {"nodes": None} must not raise.
    assert _resolve_model_id_from_graph({"nodes": None, "edges": None}) is None


def test_resolve_model_id_prefers_edge_connected_llm():
    """With multiple LLM nodes carrying DIFFERENT models, resolve the one EDGE-CONNECTED to the
    extractor (the model that actually drives extraction), not merely the first node."""
    from agentcore.services.idp.agent_config import _resolve_model_id_from_graph

    model_a, model_b = str(uuid4()), str(uuid4())

    def _llm(node_id, mid):
        return {"id": node_id, "type": "genericNode",
                "data": {"node": {"display_name": "Large Language Model",
                                  "template": {"registry_model": {"value": f"G | m | {mid}"}}}}}

    ext = {"id": "ext1", "type": "genericNode",
           "data": {"node": {"display_name": "AI Field Extractor", "template": {"config_name": {"value": ""}}}}}
    data = {
        "nodes": [_llm("llmA", model_a), _llm("llmB", model_b), ext],   # A is first by order
        "edges": [{"source": "llmB", "target": "ext1"}],                # ...but B feeds the extractor
    }
    assert _resolve_model_id_from_graph(data) == model_b   # edge-connected wins over node order


@pytest.mark.anyio
async def test_named_config_inferred_when_config_name_set_without_mode():
    """An AI Field Extractor node that has a Field Configuration selected but NO extraction_mode
    field (the universal pipeline node) must resolve to named_config — a chosen config is never
    silently ignored. (Universal-agent blocker fix.)"""
    from agentcore.services.database.models.agent.model import Agent
    from agentcore.services.database.models.idp.config import IdpAgent, IdpFieldConfiguration, IdpFieldConfigHeader
    from agentcore.services.deps import session_scope

    name = f"InferTest-{uuid4().hex[:8]}"
    config_id = uuid4()
    async with session_scope() as session:
        session.add(IdpFieldConfiguration(id=config_id, name=name, is_active=True, is_template=False))
        await session.flush()
        session.add(IdpFieldConfigHeader(id=uuid4(), config_id=config_id, field_name="total", field_type="number", display_order=0))
        await session.commit()

    try:
        # extractor node carries ONLY config_name (no extraction_mode key) — exactly the universal template.
        graph = {
            "nodes": [
                _node("AI Field Extractor", {"config_name": name}),
                _node("Large Language Model", {"registry_model": f"Groq | m | {uuid4()}"}),
            ],
            "edges": [],
        }
        # IdpAgent default mode is dynamic — the node's config selection must still win.
        idp_agent = IdpAgent(id=uuid4(), agent_id=uuid4(), extraction_mode="dynamic_prompting")
        base_agent = Agent(id=uuid4(), name="x", data=graph)

        async with session_scope() as session:
            cfg = await resolve_pipeline_config(session, idp_agent, base_agent)
        assert cfg.extraction_mode == "named_config"      # inferred from config_name
        assert cfg.field_config_id == config_id           # config actually resolved
    finally:
        async with session_scope() as session:
            row = await session.get(IdpFieldConfiguration, config_id)
            if row:
                await session.delete(row)
            await session.commit()
